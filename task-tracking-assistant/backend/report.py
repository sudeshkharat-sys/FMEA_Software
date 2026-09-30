from datetime import date, datetime, timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .db import conn

HEAD = PatternFill("solid", fgColor="1F4E78")
COLS = ["Task", "Priority", "Deadline", "Status", "Progress %", "Blocker", "Remarks", "Last updated"]


def _style_header(ws, row=1):
    for cell in ws[row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def _widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build(period: str) -> tuple[bytes, str]:
    days = 7 if period == "weekly" else 30
    since = (datetime.now() - timedelta(days=days)).isoformat(timespec="seconds")
    today = date.today().isoformat()
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary.append([f"{period.capitalize()} report", f"{(date.today() - timedelta(days=days)).isoformat()} to {today}"])
    summary["A1"].font = Font(bold=True, size=14)
    summary.append([])
    summary.append(["Employee", "Tasks", "Done", "In Progress", "Blocked", "Overdue", "Avg progress %", "Open blockers"])
    _style_header(summary, 3)

    with conn() as c:
        emps = c.execute("SELECT * FROM users WHERE role='employee' ORDER BY id").fetchall()
        for e in emps:
            # tasks touched in the period, or still open
            tasks = c.execute(
                "SELECT * FROM tasks WHERE employee_id=? AND (updated_at>=? OR status!='Done') ORDER BY deadline",
                (e["id"], since),
            ).fetchall()
            n = len(tasks)
            done = sum(t["status"] == "Done" for t in tasks)
            prog = sum(t["status"] == "In Progress" for t in tasks)
            blocked = sum(t["status"] == "Blocked" for t in tasks)
            overdue = sum(bool(t["deadline"]) and t["deadline"] < today and t["status"] != "Done" for t in tasks)
            avg = round(sum(t["progress"] for t in tasks) / n) if n else 0
            blockers = "; ".join(t["blocker"] for t in tasks if t["blocker"] and t["status"] != "Done")
            summary.append([e["name"], n, done, prog, blocked, overdue, avg, blockers])

            ws = wb.create_sheet(e["name"][:31])
            ws.append(COLS)
            _style_header(ws)
            for t in tasks:
                ws.append([t["title"], t["priority"], t["deadline"], t["status"], t["progress"],
                           t["blocker"], t["remarks"], t["updated_at"].replace("T", " ")])
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
            _widths(ws, [38, 10, 12, 13, 11, 34, 34, 18])
            ws.freeze_panes = "A2"
    _widths(summary, [16, 8, 8, 12, 9, 9, 15, 60])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue(), f"{period}_report_{today}.xlsx"
